import os; clear = lambda: os.system('cls'); clear();
import openpyxl as xl; #print("Python Library: ", xl.__name__, xl.__version__)
import matplotlib.pyplot as plt; from openpyxl.drawing.image import Image; import seaborn as sns
import pandas as pd; import warnings;
warnings.filterwarnings("ignore", category=RuntimeWarning)

################################################################################

path_in = r'C:\Users\jordi\Desktop\ARCHIVO\1 - DOCS\PyBcn - OpenPyXL\XLS\ex1_countries.xlsx'
path_out = r'C:\Users\jordi\Desktop\ARCHIVO\1 - DOCS\PyBcn - OpenPyXL\XLS\ex3_plots.xlsx'

################################################################################

print("\n***************************************************************************")
print("DB: Northwind; TB: OrderID")
print("***************************************************************************\n")

df = pd.read_excel(path_in, sheet_name=None)

wb = xl.Workbook()
wb.remove(wb.active)

def add_charts_sheet(wb, sheet_name, data):

    ws = wb.create_sheet(title=sheet_name)
    plt.figure(figsize=(15, 10))
    plt.subplot(2, 2, 1)
    data['ShipCity'].value_counts().plot(kind='bar', title='Frecuencia de ShipCity')
    plt.savefig('bar_chart.png')
    plt.clf()
    img1 = Image('bar_chart.png')
    ws.add_image(img1, 'A1')

for sheet_name, data in df.items():
    add_charts_sheet(wb, sheet_name + ' Graphs', data)

wb.save(path_out)

os.remove(r"C:\Users\jordi\Desktop\ARCHIVO\1 - DOCS\PyBcn - OpenPyXL\bar_chart.png")

print("\nOK 😊👍\n")

print("\n***************************************************************************\n")

################################################################################
